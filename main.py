import json
import os
import io
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv

from ai_extractor import extract_from_text, extract_from_image, chat_response
from form_schema import VALID_FIELDS, REQUIRED_FIELDS, FIELD_LABELS

load_dotenv()

app = FastAPI(
    title="EHR Form AI API",
    description="AI-powered blood bank request form assistant with document/image extraction",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_api_key() -> str:
    key = os.getenv("GEMINI_API_KEY", "")
    if not key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY not configured")
    return key


# ---------- Document Upload Endpoint ----------

@app.post("/api/extract", summary="Extract form data from uploaded file")
async def extract_from_file(
    file: UploadFile = File(...),
    currentForm: str = Form(default="{}"),
):
    """
    Upload a PDF, DOCX, or image file to extract blood bank request form data.

    - **file**: The document or image file (PDF, DOCX, PNG, JPG, JPEG)
    - **currentForm**: JSON string of current form state (optional)

    Returns extracted field values, missing required fields, and a summary.
    """
    api_key = get_api_key()
    file_bytes = await file.read()

    try:
        current_form = json.loads(currentForm)
    except json.JSONDecodeError:
        current_form = {}

    content_type = file.content_type or ""
    filename = (file.filename or "").lower()

    # Image files
    if content_type.startswith("image/") or filename.endswith((".png", ".jpg", ".jpeg", ".webp")):
        mime = content_type if content_type.startswith("image/") else "image/png"
        result = await extract_from_image(file_bytes, current_form, api_key, mime)
        return result

    # PDF
    if content_type == "application/pdf" or filename.endswith(".pdf"):
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            raise HTTPException(status_code=400, detail="Failed to parse PDF file")

        if not text.strip():
            raise HTTPException(status_code=400, detail="Could not extract text from PDF. Try uploading as an image.")

        result = await extract_from_text(text, current_form, api_key)
        return result

    # DOCX
    if content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) or filename.endswith((".docx", ".doc")):
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))

            # Extract text from paragraphs
            parts = [para.text for para in doc.paragraphs if para.text.strip()]

            # Extract text from tables (blood request forms often use tables)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))

            text = "\n".join(parts)
        except Exception:
            raise HTTPException(status_code=400, detail="Failed to parse DOCX file")

        if not text.strip():
            raise HTTPException(status_code=400, detail="Could not extract text from document.")

        result = await extract_from_text(text, current_form, api_key)
        return result

    raise HTTPException(
        status_code=400,
        detail=f"Unsupported file type: {content_type}. Supported: PDF, DOCX, PNG, JPG, JPEG",
    )


# ---------- Chat Assistant Endpoint ----------

class ChatRequest(BaseModel):
    user: str
    currentForm: dict = {}
    pending: Optional[str] = None
    isQuestion: bool = False


@app.post("/api/chat", summary="Chat with the form assistant")
async def chat_assistant(req: ChatRequest):
    """
    Conversational form-filling assistant.

    - **user**: The user's message
    - **currentForm**: Current form state as JSON object
    - **pending**: The field currently being asked about
    - **isQuestion**: If true, treats the message as a question only

    Returns a reply, field updates, and optional submit flag.
    """
    if not req.user.strip():
        raise HTTPException(status_code=400, detail="Message cannot be empty")

    api_key = get_api_key()
    result = await chat_response(
        user_message=req.user,
        current_form=req.currentForm,
        pending_field=req.pending,
        is_question=req.isQuestion,
        api_key=api_key,
    )
    return result


# ---------- Form Schema Endpoint ----------

@app.get("/api/schema", summary="Get form field schema")
async def get_schema():
    """Returns the form field definitions, labels, and required fields."""
    return {
        "fields": VALID_FIELDS,
        "required": REQUIRED_FIELDS,
        "labels": FIELD_LABELS,
    }


# ---------- Health Check ----------

@app.get("/api/health")
async def health():
    has_key = bool(os.getenv("GEMINI_API_KEY"))
    return {"status": "ok", "gemini_configured": has_key}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
